"""Requirements ingestion: file -> pages -> chunks -> ChromaDB.

Contract function:
    ingest_requirements(pdf_path) -> int      # number of chunks stored

Accepts a path (str/Path), raw bytes, or a Streamlit UploadedFile.
Supports PDF (main), DOCX and TXT. This module reads PDFs directly with PyMuPDF so it
does not depend on Member 2's utils/pdf_reader.py.

CLI:  python -m rag.ingest data/requirements/my_requirements.pdf
"""
from __future__ import annotations

import io
import logging
import sys
from pathlib import Path
from typing import Any, Optional

from rag import vectorstore
from rag.chunker import MAX_CHARS, chunk_pages

try:
    import pymupdf as fitz  # PyMuPDF >= 1.24
except ImportError:  # pragma: no cover
    import fitz  # type: ignore

log = logging.getLogger(__name__)
MIN_TEXT_CHARS = 30


# ------------------------------- reading ----------------------------------- #
def _resolve_input(source: Any, filename: Optional[str]) -> tuple[Optional[bytes], str]:
    """Return (bytes or None if a path, filename)."""
    if isinstance(source, (str, Path)):
        return None, filename or str(source)
    if hasattr(source, "getvalue"):  # Streamlit UploadedFile / BytesIO
        return source.getvalue(), filename or getattr(source, "name", "upload.pdf")
    if hasattr(source, "read"):
        return source.read(), filename or getattr(source, "name", "upload.pdf")
    if isinstance(source, (bytes, bytearray)):
        return bytes(source), filename or "upload.pdf"
    raise TypeError(f"Unsupported source type: {type(source)!r}")


def _read_pdf(data: Optional[bytes], path: str) -> list[tuple[int, str]]:
    doc = fitz.open(stream=data, filetype="pdf") if data is not None else fitz.open(path)
    try:
        return [(i + 1, page.get_text("text", sort=True)) for i, page in enumerate(doc)]
    finally:
        doc.close()


def _read_docx(data: Optional[bytes], path: str) -> list[tuple[int, str]]:
    from docx import Document  # lazy import

    doc = Document(io.BytesIO(data)) if data is not None else Document(path)
    lines: list[str] = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            lines.append("")
        elif (p.style.name or "").lower().startswith("heading"):
            lines.append(f"## {txt}")
        else:
            lines.append(txt)
    for table in doc.tables:
        for row in table.rows:
            lines.append("- " + " | ".join(c.text.strip() for c in row.cells))
    return [(1, "\n".join(lines))]


def _read_txt(data: Optional[bytes], path: str) -> list[tuple[int, str]]:
    raw = data if data is not None else Path(path).read_bytes()
    text = raw.decode("utf-8", errors="replace")
    return [(i + 1, t) for i, t in enumerate(text.split("\f"))]


def read_requirements_pages(source: Any, filename: Optional[str] = None) -> list[tuple[int, str]]:
    """Return [(page_number, text)] for a PDF / DOCX / TXT requirements file."""
    data, name = _resolve_input(source, filename)
    ext = Path(name).suffix.lower()
    if ext == ".docx":
        pages = _read_docx(data, name)
    elif ext in (".txt", ".md"):
        pages = _read_txt(data, name)
    else:  # default: PDF
        pages = _read_pdf(data, name)
    if sum(len(t.strip()) for _, t in pages) < MIN_TEXT_CHARS:
        raise ValueError(
            "No extractable text found in the requirements file. It may be a scanned PDF. "
            "Use a text-based PDF for the requirements document."
        )
    return pages


def _default_source_name(name: str) -> str:
    return Path(name).stem.replace("_", " ").replace("-", " ").strip().title() or "Requirements"


# ------------------------------- ingestion --------------------------------- #
def ingest_pages(pages, source_name: str, reset: bool = True, max_chars: int = MAX_CHARS) -> int:
    """Chunk already-extracted pages and store them. Returns chunk count."""
    chunks = chunk_pages(pages, source=source_name, max_chars=max_chars)
    if not chunks:
        raise ValueError("Requirements text produced zero chunks.")
    if reset:
        vectorstore.reset_collection()
    n = vectorstore.add_chunks(chunks)
    log.info("Ingested %d chunks from %s", n, source_name)
    return n


def ingest_requirements(
    source: Any,
    source_name: Optional[str] = None,
    reset: bool = True,
    filename: Optional[str] = None,
) -> int:
    """Contract function. Read the requirements file, chunk it, store it in Chroma.

    reset=True (default) wipes any previous requirements so each application starts clean.
    """
    pages = read_requirements_pages(source, filename)
    _, name = _resolve_input(source, filename)
    return ingest_pages(pages, source_name or _default_source_name(name), reset=reset)


if __name__ == "__main__":  # pragma: no cover
    if len(sys.argv) < 2:
        sys.exit("usage: python -m rag.ingest <requirements.pdf>")
    count = ingest_requirements(sys.argv[1])
    print(f"Stored {count} chunks in {vectorstore.chroma_dir()}")
